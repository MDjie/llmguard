"""Create synthetic codec fixtures, never semantic quality or independent holdout data."""
from pathlib import Path
import sys
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

output = Path(sys.argv[1]); output.mkdir(parents=True, exist_ok=True)
first = Image.new('RGB', (160, 100), 'white')
ImageDraw.Draw(first).text((8, 20), 'FORMAT FIXTURE', fill='black')
second = Image.new('RGB', (160, 100), 'yellow')
for extension, kind in [('png', 'PNG'), ('jpg', 'JPEG'), ('webp', 'WEBP'), ('bmp', 'BMP')]:
    first.save(output / ('sample.' + extension), format=kind)
first.save(output / 'sample.gif', save_all=True, append_images=[second], duration=100, loop=0)
first.save(output / 'sample.tiff', save_all=True, append_images=[second])
first.save(output / 'sample.pdf', save_all=True, append_images=[second])
document = Document(); document.add_paragraph('DOCUMENT FORMAT FIXTURE')
document.save(output / 'sample.docx')
workbook = Workbook(); workbook.active['A1'] = 'VISIBLE CONTENT'
hidden = workbook.create_sheet('hidden'); hidden['A1'] = 'HIDDEN CONTENT MUST BE INSPECTED'; hidden.sheet_state = 'hidden'
workbook.save(output / 'sample.xlsx')
slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text = 'FORMAT FIXTURE'
slide.notes_slide.notes_text_frame.text = 'SPEAKER NOTES MUST BE INSPECTED'
slides.save(output / 'sample.pptx')
(output / 'sample.rtf').write_bytes(b'{\\rtf1\\ansi FORMAT FIXTURE}')
(output / 'bom-utf16.txt').write_bytes('安全文本'.encode('utf-16'))
(output / 'legacy-gb18030.txt').write_bytes('安全文本'.encode('gb18030'))

# Standalone subtitles retain the original UTF-16 character-coordinate and cue timeline.
(output / 'sample.srt').write_bytes('1\r\n00:00:01,000 --> 00:00:03,000\r\n欢迎参加产品培训。\r\n\r\n2\r\n00:00:04,000 --> 00:00:06,000\r\n今天讨论服务流程。\r\n'.encode('utf-8'))
(output / 'sample.vtt').write_bytes('WEBVTT\n\n00:01.000 --> 00:03.000\n欢迎参加产品培训。\n\n00:04.000 --> 00:06.000\n今天讨论服务流程。\n'.encode('utf-8'))
(output / 'sample.ass').write_bytes('[Script Info]\nScriptType: v4.00+\n[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\nDialogue: 0,0:00:01.00,0:00:03.00,Default,,0,0,0,,欢迎参加产品培训。\n'.encode('utf-8'))
(output / 'sample.ssa').write_bytes('[Script Info]\nScriptType: v4.00\n[Events]\nFormat: Marked, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\nDialogue: Marked=0,0:00:01.00,0:00:03.00,Default,,0,0,0,,欢迎参加产品培训。\n'.encode('utf-8'))
