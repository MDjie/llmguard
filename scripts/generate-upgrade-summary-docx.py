# -*- coding: utf-8 -*-
"""Generate the GuardLLM major-release summary as a polished DOCX."""

from __future__ import annotations

import argparse
import math
import tempfile
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "docs" / "国舜大模型安全护栏大版本升级总结_2026-09-02.docx"
LOGO_PATH = ROOT / "public" / "logo.png"

PAGE_WIDTH_DXA = 12_240
PAGE_HEIGHT_DXA = 15_840
CONTENT_WIDTH_DXA = 9_360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"

BLUE = "2E74B5"
DARK_BLUE = "17324D"
DEEP_BLUE = "1F4D78"
TEAL = "1D7A6B"
GOLD = "C9922E"
BLACK = "1A1A1A"
GRAY = "5B6573"
MUTED = "7A8491"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F4F1"
PALE_GOLD = "FBF3E4"
BORDER = "D8DEE8"
WHITE = "FFFFFF"


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.2,
                  color: str = BLACK, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def set_run_font(run, *, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 latin: str = LATIN_FONT, east_asia: str = CJK_FONT) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, size: float, color: str = BLACK, bold: bool = False,
                   italic: bool = False) -> None:
    style.font.name = LATIN_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, size=11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = doc.styles["Title"]
    set_style_font(title, size=30, color=DARK_BLUE, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, size=14, color=GRAY)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DEEP_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    set_style_font(caption, size=9, color=MUTED, italic=True)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    if "Table Citation" not in [style.name for style in doc.styles]:
        citation = doc.styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        citation = doc.styles["Table Citation"]
    set_style_font(citation, size=9, color=MUTED, italic=True)
    citation.paragraph_format.space_before = Pt(4)
    citation.paragraph_format.space_after = Pt(4)

    if "Code Reference" not in [style.name for style in doc.styles]:
        code = doc.styles.add_style("Code Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code Reference"]
    set_style_font(code, size=9, color=GRAY)
    code._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), MONO_FONT)
    code._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), MONO_FONT)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(3)
    code.paragraph_format.left_indent = Inches(0.22)


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def add_page_field(paragraph) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)
    set_run_font(run, size=8.5, color=MUTED)


def configure_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("GuardLLM 大版本升级总结")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = p.add_run("\t代码基线 main@41e3729")
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    prefix = fp.add_run("内部版本总结  ·  第 ")
    set_run_font(prefix, size=8.5, color=MUTED)
    add_page_field(fp)
    suffix = fp.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = ""


def add_numbering_definition(doc: Document, *, kind: str, abstract_id: int,
                             num_id: int) -> int:
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(lvl_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    ppr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    if kind == "bullet":
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        rpr.append(fonts)
        level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_list_item(doc: Document, text: str, num_id: int, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=11, color=BLACK, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color=BLACK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=BLACK)
    return p


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in CELL_MARGINS_DXA.items():
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_table_borders(table, *, color: str = BORDER, size: int = 4,
                      left_color: str | None = None, left_size: int | None = None) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(left_size if edge == "left" and left_size else size))
        node.set(qn("w:color"), left_color if edge == "left" and left_color else color)
        node.set(qn("w:space"), "0")


def set_table_geometry(table, widths: list[int], *, indent: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int],
              *, aligns: list | None = None, font_size: float = 9.2,
              status_col: int | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        shade_cell(header.cells[index], LIGHT_GRAY)
        set_cell_text(
            header.cells[index], label, bold=True, size=9.2, color=DARK_BLUE,
            align=(aligns[index] if aligns else WD_ALIGN_PARAGRAPH.CENTER),
        )
    for values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for index, value in enumerate(values):
            alignment = aligns[index] if aligns else WD_ALIGN_PARAGRAPH.LEFT
            color = BLACK
            bold = False
            if status_col is not None and index == status_col:
                bold = True
                if "完成" in value or "通过" in value:
                    color = TEAL
                elif "待" in value or "部分" in value or "目标" in value:
                    color = GOLD
            set_cell_text(row.cells[index], value, bold=bold, size=font_size,
                          color=color, align=alignment)
    set_table_geometry(table, widths)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc: Document, label: str, text: str, *, fill: str = PALE_BLUE,
                accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    lead = p.add_run(f"{label}  ")
    set_run_font(lead, size=10.5, color=accent, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=10.5, color=BLACK)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=fill, size=0, left_color=accent, left_size=18)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def add_metric_strip(doc: Document) -> None:
    metrics = [
        ("634", "变更文件", PALE_BLUE, BLUE),
        ("+58,348 / -8,148", "代码变更", PALE_TEAL, TEAL),
        ("300", "自动化用例", PALE_GOLD, GOLD),
        ("29 / 12", "迁移 / 关键关系", LIGHT_GRAY, DARK_BLUE),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_repeat_table_header(table.rows[0])
    for index, (value, label, fill, accent) in enumerate(metrics):
        cell = table.cell(0, index)
        shade_cell(cell, fill)
        cell.text = ""
        value_p = cell.paragraphs[0]
        value_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_p.paragraph_format.space_after = Pt(2)
        value_run = value_p.add_run(value)
        set_run_font(value_run, size=15 if index != 1 else 12, color=accent, bold=True)
        label_p = cell.add_paragraph()
        label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_p.paragraph_format.space_before = Pt(0)
        label_p.paragraph_format.space_after = Pt(2)
        label_run = label_p.add_run(label)
        set_run_font(label_run, size=8.8, color=GRAY, bold=True)
    set_table_geometry(table, [2340, 2340, 2340, 2340])
    set_table_borders(table, color=WHITE, size=8)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_label_paragraph(doc: Document, label: str, value: str, *, after: float = 3) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    label_run = p.add_run(f"{label}  ")
    set_run_font(label_run, size=10.5, color=MUTED, bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, size=10.5, color=BLACK)


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def load_font(size: int, *, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
                       font, fill: str, *, line_gap: int = 6) -> None:
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bounds = draw.textbbox((0, 0), line, font=font)
        widths.append(bounds[2] - bounds[0])
        heights.append(bounds[3] - bounds[1])
    total_height = sum(heights) + line_gap * (len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total_height) / 2
    for line, width, height in zip(lines, widths, heights):
        x = box[0] + (box[2] - box[0] - width) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + line_gap


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
               color: str = "6B7785", width: int = 5) -> None:
    normalized_color = color if color.startswith("#") else f"#{color}"
    draw.line([start, end], fill=normalized_color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 16
    spread = 0.55
    p1 = (end[0] - length * math.cos(angle - spread),
          end[1] - length * math.sin(angle - spread))
    p2 = (end[0] - length * math.cos(angle + spread),
          end[1] - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=normalized_color)


def create_architecture_image(path: Path) -> None:
    image = Image.new("RGB", (1800, 900), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    title_font = load_font(42, bold=True)
    box_font = load_font(29, bold=True)
    small_font = load_font(24)
    draw.text((90, 50), "升级后的 GuardLLM 安全数据面", font=title_font, fill="#17324D")
    draw.text((90, 112), "统一入口、统一证据、统一决策、统一审计", font=small_font, fill="#5B6573")

    top_boxes = [
        ((90, 210, 380, 360), "业务应用\nSDK / Agent", PALE_BLUE, BLUE),
        ((500, 210, 840, 360), "Guard Gateway\nREST / SSE / WS / gRPC", PALE_TEAL, TEAL),
        ((970, 210, 1280, 360), "GuardEngine V2\n策略与统一决策", PALE_GOLD, GOLD),
        ((1410, 210, 1710, 360), "模型与内容交付\n输出复检 / 标识", LIGHT_GRAY, DARK_BLUE),
    ]
    for box, text, fill, outline in top_boxes:
        draw.rounded_rectangle(box, radius=20, fill=f"#{fill}", outline=f"#{outline}", width=4)
        draw_centered_text(draw, box, text, box_font, f"#{outline}")
    draw_arrow(draw, (380, 285), (500, 285))
    draw_arrow(draw, (840, 285), (970, 285))
    draw_arrow(draw, (1280, 285), (1410, 285))

    modules = [
        ((140, 455, 500, 590), "文本 / DLP\n推理攻击与多轮累计", PALE_BLUE, BLUE),
        ((550, 455, 910, 590), "OCR / VLM / ASR\n视频时间轴与跨模态融合", PALE_TEAL, TEAL),
        ((960, 455, 1320, 590), "RAG 来源可信\n工具授权与结果复检", PALE_GOLD, GOLD),
        ((1370, 455, 1660, 590), "Judge / 内容标识\n事件与人工复核", LIGHT_GRAY, DARK_BLUE),
    ]
    for box, text, fill, outline in modules:
        draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{outline}", width=3)
        draw_centered_text(draw, box, text, small_font, f"#{outline}")
        draw.line([(1125, 360), ((box[0] + box[2]) // 2, box[1])], fill="#9AA5B1", width=3)

    bottom = (90, 700, 1710, 820)
    draw.rounded_rectangle(bottom, radius=18, fill="#F7F9FB", outline="#D8DEE8", width=3)
    bottom_text = "租户 / 应用 / 主体隔离    |    审计链与可信时间    |    数据谱系与删除证明    |    可观测性与安全运营"
    draw_centered_text(draw, bottom, bottom_text, small_font, "#46515F")
    for x in (320, 720, 1120, 1520):
        draw_arrow(draw, (x, 700), (x, 610), color="9AA5B1", width=3)
    image.save(path, format="PNG", optimize=True)


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_cover(doc: Document) -> None:
    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        shape = p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.55))
        set_picture_alt_text(shape, "国舜科技标识", "国舜大模型安全护栏品牌标识")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("GUARDLLM  ·  MAJOR RELEASE REVIEW")
    set_run_font(run, size=10.5, color=GOLD, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("国舜大模型安全护栏\n大版本升级总结")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("功能、性能与场景价值说明")

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.space_before = Pt(8)
    lead.paragraph_format.space_after = Pt(44)
    run = lead.add_run("从规则检测平台升级为面向企业生产链路的统一 AI 安全治理底座")
    set_run_font(run, size=11, color=GRAY, italic=True)

    for label, value in [
        ("版本范围", "af45893  →  41e3729"),
        ("编制日期", "2026-09-02"),
        ("文档版本", "V1.0"),
        ("适用对象", "产品、研发、交付、售前与安全运营团队"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        label_run = p.add_run(f"{label}  ")
        set_run_font(label_run, size=9.5, color=MUTED, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=9.5, color=BLACK)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(35)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(0)
    note_run = note.add_run("内部版本总结  |  结论以代码、自动化测试与目标环境验收证据分层表述")
    set_run_font(note_run, size=8.5, color=MUTED)
    add_page_break(doc)


def add_executive_summary(doc: Document, bullet_id: int) -> None:
    doc.add_heading("执行摘要", level=1)
    add_callout(
        doc,
        "核心结论",
        "本次升级把 GuardLLM 从以控制台和规则检测为主的应用，推进为覆盖统一安全接入、"
        "GuardEngine V2、多模态异步分析、RAG/Agent 信任边界、策略发布、审计与生产化部署的企业级平台。"
        "代码内不存在完全无映射的需求，但目标模型效果、真实环境性能、企业系统联调和客户签名验收仍需后续完成。",
        fill=PALE_BLUE,
        accent=BLUE,
    )
    add_metric_strip(doc)
    doc.add_heading("本次升级带来的四个变化", level=2)
    points = [
        "安全能力从单次文本规则命中扩展到有界归一化、多轮会话、分步推理攻击、推理过程提取与跨模态联合证据。",
        "接入形态从单一 Web/API 扩展到 REST、SSE、WebSocket、gRPC 与 Java/Python/Go/TypeScript SDK，并通过统一契约维持兼容。",
        "处理链路从同步小文件扩展到分片上传、异步任务、OCR/VLM/ASR、音视频时间轴和内容标识，支持大文件的受控分析。",
        "治理方式从功能可用扩展到租户隔离、审计链、可信时间、数据谱系、事件闭环、Helm 高可用模板与可追溯验收框架。",
    ]
    for point in points:
        add_list_item(doc, point, bullet_id)

    doc.add_heading("完成度口径", level=2)
    add_table(
        doc,
        ["状态", "数量", "含义"],
        [
            ["代码完成、待目标验收", "22", "代码及本地自动化具备，仍需真实依赖或目标环境 POC"],
            ["部分实现", "64", "已有实现路径，但仍缺功能、模型校准、数据或外部集成"],
            ["仅验收框架", "15", "性能、可用性与安全测评脚本具备，尚无目标环境报告"],
            ["完全未实现", "0", "101 项均存在代码、测试或验收框架映射"],
            ["正式签名验收", "0 / 101", "尚未取得客户签名证据，不应表述为正式验收通过"],
        ],
        [2200, 1200, 5960],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        status_col=0,
    )
def add_scope_and_architecture(doc: Document, bullet_id: int, architecture_path: Path) -> None:
    add_page_break(doc)
    doc.add_heading("1. 升级范围与总体变化", level=1)
    add_label_paragraph(doc, "升级基线", "af45893（维度管理完善 + 敏感词库 / NeMo 规则导入工具）")
    add_label_paragraph(doc, "平台升级", "6790142（企业级安全护栏平台）")
    add_label_paragraph(doc, "检测增强", "41e3729（多模态与推理攻击检测）")
    add_label_paragraph(doc, "累计规模", "634 个文件变更，新增 58,348 行，删除 8,148 行")

    doc.add_heading("升级前后对比", level=2)
    add_table(
        doc,
        ["维度", "升级前", "升级后"],
        [
            ["检测范围", "文本规则、关键词、基础 Judge", "有界归一化、推理攻击、多轮会话、多模态与 RAG/Agent 联合检测"],
            ["接入方式", "Web 控制台与基础 API", "REST / SSE / WebSocket / gRPC、旁路与四语言 SDK"],
            ["媒体处理", "文档/图片处理能力分散", "分片上传、异步任务、多视图 OCR/VLM/ASR、音视频时间轴"],
            ["策略治理", "策略编辑与基础测试", "签名策略包、审批分离、影子/灰度/激活/回滚、异步评测"],
            ["数据安全", "基础登录与数据表", "租户/应用/主体隔离、审计链、可信时间、谱系与删除证明"],
            ["生产交付", "单体部署为主", "Helm/Kustomize、网络隔离、mTLS、PDB/HPA、离线包与可观测性"],
        ],
        [1540, 3100, 4720],
        font_size=9.1,
    )

    doc.add_heading("目标架构", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(architecture_path), width=Inches(6.35))
    set_picture_alt_text(
        shape,
        "GuardLLM 升级架构",
        "业务应用通过 Guard Gateway 进入 GuardEngine V2，联动文本、推理、多模态、RAG 和工具安全模块，"
        "并由租户隔离、审计链、数据谱系和可观测性提供底层保障。",
    )
    caption = doc.add_paragraph(style="Caption")
    caption.add_run("图 1  升级后的统一安全数据面与治理底座")
    add_callout(
        doc,
        "安全边界",
        "模型调用前必须完成身份上下文校验、配额/并发控制和输入检测；流式内容在提交客户端前经过输出检测；"
        "必需检测器按策略 fail-closed，原始敏感内容不复制到安全事件。",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def add_feature_sections(doc: Document, bullet_id: int) -> None:
    add_page_break(doc)
    doc.add_heading("2. 功能升级清单", level=1)
    intro = doc.add_paragraph(
        "以下清单按产品能力域归纳。标注“已完成”的项目表示代码与本地自动化已具备；涉及真实模型、企业平台或"
        "客户指标的项目仍需目标环境验收。"
    )
    intro.paragraph_format.space_after = Pt(8)

    sections = [
        (
            "2.1 统一安全接入与跨语言契约",
            [
                "安全网关支持 OpenAI 兼容 REST、SSE、WebSocket 与 gRPC，统一输入/输出检测、Trace 与错误语义。",
                "请求上下文签名绑定租户、应用、主体、凭证、请求、会话和绝对截止时间，并支持 API Key/JWT/mTLS 边界。",
                "支持模型权重路由、JSON Pointer 请求/响应字段映射、思考字段归一化和流式提交门。",
                "Guard v1 由单一 Schema 生成 TypeScript、Java、Python、OpenAPI 与 Protobuf，CI 检查生成漂移和向后兼容。",
                "提供 Java、Python、Go、TypeScript SDK；动态模型路由控制台和企业统一身份联调仍属于后续范围。",
            ],
        ),
        (
            "2.2 GuardEngine V2 与推理攻击检测",
            [
                "有界归一化覆盖 NFKC、零宽字符、URL/HTML 实体、Base64/Base32/Hex、转义码点、Quoted-Printable、ROT13 与同形字符。",
                "新增分步推理攻击检测：识别步骤序列、绕过方法、受保护目标、隐蔽手法和有害目标，并按多轮推进累计风险。",
                "新增推理过程提取检测：识别中英文思维链、内部推理与草稿泄露请求，生成掩码预览和 HMAC 证据。",
                "会话上下文不再因中间轮次提前退出而丢失后续证据，保留匿名步骤标识、累计分和最终阻断点。",
                "内置 DLP、资源消耗攻击、保险收益保证/虚假理赔承诺/健康告知规避/条款弱化/高压营销等基线。",
            ],
        ),
        (
            "2.3 图片、文档、音频与视频安全",
            [
                "图片/文档通过旋转、阈值、去噪等视图并行执行 OCR 和视觉分类，再将文本、视觉和用户上下文联合决策。",
                "音频新增原始、降噪、响度归一化、0.9/1.1 倍速和反向探测视图；多视图 ASR 结果完成时间戳回映、去重和置信度过滤。",
                "视频按固定间隔或时长自适应抽帧，对帧执行 OCR/视觉分析，并与音轨 ASR、声学异常和用户文本统一到时间轴。",
                "跨模态融合按可配置时间窗口聚类；单模态低风险但联合意图升级时，输出 cooperativeAttack 及完整来源证据。",
                "媒体分析器采用外部命令适配器接入 OCR/VLM/ASR/音频分类模型；生产召回率取决于目标模型与数据集，需 POC。",
            ],
        ),
        (
            "2.4 RAG 与 Agent / 工具安全",
            [
                "知识入库记录来源、版本、切片、ACL、污染标记和扫描结论，检索结果可追溯原文档与片段。",
                "区分系统指令、用户输入、检索上下文与工具返回的信任等级，阻止低信任内容覆盖高优先级安全规则。",
                "工具调用执行白名单、参数 Schema、资源范围、频率与最小权限校验，高风险动作支持人工/双人审批。",
                "工具返回结果再次进入检测链，降低间接注入、数据外传、凭证泄露与越权执行风险。",
            ],
        ),
        (
            "2.5 策略发布、评测与运营闭环",
            [
                "策略包采用确定性编译、内容哈希与签名校验，支持草稿、测试、审批、影子、灰度、激活、回滚、撤回和归档。",
                "发布控制台强制创建者与审批者分离并使用乐观锁，降低误发布和越权发布风险。",
                "异步评测支持幂等提交、队列进度、失败重试、准确率/召回率/FPR/FNR/F1 与 badcase 查看。",
                "安全事件支持责任人、SLA、状态机、证据详情与流转时间线；误报反馈到样本/规则/发布的闭环仍需完善。",
            ],
        ),
        (
            "2.6 身份、租户、数据与审计",
            [
                "API 安全边界统一实施认证授权、CSRF、请求体限制、限流、响应 Schema、审计和安全错误返回。",
                "数据库查询按租户/应用约束；密码策略、失败锁定、历史密码、过期与权限矩阵纳入自动化。",
                "审计事件具备前序哈希、签名、追加写、独立验证与企业可信时间适配。",
                "谱系覆盖模型输出、内容标识、媒体衍生物、文档图像和 RAG 来源/切片；清理生成 HMAC 删除证明。",
                "企业 IdP/MFA、KMS/TSA、组织树和三员分立仍需客户平台联调。",
            ],
        ),
        (
            "2.7 内容标识与开放治理",
            [
                "内容标识覆盖文本、图片、音频、视频的显式/隐式标识和一致性验证。",
                "内容标识与输出检测、媒体衍生物和谱系协同，支持传播追踪与异常告警。",
                "策略、告警、审计、评测与报表已形成开放 API 基础，完整联调仍需目标环境确认。",
            ],
        ),
        (
            "2.8 部署、供应链与可观测性",
            [
                "Helm 将网关、控制面、媒体、推理、数据访问与可观测组件分域部署，提供 NetworkPolicy、mTLS、PDB、HPA、最小权限及 amd64/arm64 调度抽象；昇腾配置为模板而非互认证结论。",
                "离线交付校验镜像摘要、Cosign、SBOM 与 SHA256SUMS；Prometheus、Grafana、OpenTelemetry、日志、Trace 和运行手册支撑故障定位与灾备。",
            ],
        ),
    ]
    for heading, items in sections:
        doc.add_heading(heading, level=2)
        for item in items:
            add_list_item(doc, item, bullet_id)


def add_performance_section(doc: Document, bullet_id: int) -> None:
    add_page_break(doc)
    doc.add_heading("3. 性能、容量与工程质量清单", level=1)
    add_callout(
        doc,
        "阅读口径",
        "“实现约束”表示代码已限制资源边界；“本轮实测”表示当前工作区可复现结果；“目标门槛”来自验收规范，"
        "必须在目标全策略环境取得不可变报告后才能判定通过。",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    doc.add_heading("3.1 已实现的容量与资源边界", level=2)
    add_table(
        doc,
        ["项目", "实现值 / 默认值", "产品作用"],
        [
            ["文本归一化", "最多 24 个视图；解码深度 3；单视图 1,048,576 字符", "限制编码炸弹与归一化放大，同时覆盖常见混淆变体"],
            ["会话上下文", "合同文本上限 1,048,576 字符；尾部窗口 4,096 字符；TTL 最长 7 天", "保留多轮攻击轨迹并控制内存和会话状态膨胀"],
            ["同步检测请求", "Guard API 请求体 2 MiB；绝对截止时间最长 60 秒", "防止超大同步请求拖垮在线数据面"],
            ["对象分片", "16 MiB / 片；租户默认总配额 10 GiB", "支持大文件断点与异步校验，隔离单租户资源占用"],
            ["文件类型上限", "文本/图片 100 MiB；文档/音频 500 MiB；视频 3 GiB", "将大文件转入对象存储和异步 Worker，避免 Web 进程整文件入内存"],
            ["文档与视频采样", "文档最多 5,000 页；视频最多 10,000 帧", "在覆盖率与计算成本之间提供可配置上限"],
            ["多模态批处理", "默认 4；范围 1-64；最小置信度 0.35", "控制模型并发和噪声证据进入融合链的比例"],
            ["风险策略", "复核阈值 0.65；阻断阈值 0.80；跨模态窗口 30 秒", "按业务风险偏好平衡自动阻断、人工复核与成本"],
            ["速率与并发", "租户/应用/主体/凭证/API 多级配额与 Redis 原子租约", "避免热点应用或故障客户端影响其他租户"],
        ],
        [2200, 3100, 4060],
        font_size=8.9,
    )

    doc.add_heading("3.2 本轮可复现工程结果", level=2)
    add_table(
        doc,
        ["验证项", "结果", "说明"],
        [
            ["Vitest", "76 个测试文件 / 300 个用例全部通过", "覆盖 API 安全、检测、GuardEngine V2、多模态、审计、RAG、工具、部署等模块"],
            ["整体覆盖率", "语句 27.51% / 分支 21.26% / 函数 25.38% / 行 27.72%", "覆盖率仍偏低，主要空白位于 UI、路由编排和真实子进程集成"],
            ["关键新增模块", "推理攻击行 95.31%；时序融合/检测策略/跨模态融合行 100%", "本次新增的核心决策逻辑已获得针对性高覆盖"],
            ["Media Analyzer", "独立 ESM 构建通过", "实际 OCR/VLM/ASR 命令链仍需目标模型端到端 POC"],
            ["PostgreSQL 集成", "29 个迁移；12 个关键关系；4 项检查通过", "租户隔离、谱系约束和追加审计均通过"],
            ["前后端 E2E", "桌面/移动端共 10/10 通过", "覆盖健康检查、登录、鉴权、导航、模型管理、会话与退出"],
            ["Next.js 生产构建", "通过，78 个路由完成生成/检查", "验证 App Router 与生产打包可用性"],
        ],
        [2200, 2860, 4300],
        font_size=8.9,
        status_col=1,
    )

    doc.add_heading("3.3 目标环境性能门槛（尚未签名验收）", level=2)
    targets = [
        "黑样本召回率与精准率均不低于 95%；2,000 条文本攻击拦截率不低于 99%。",
        "2,000 条白样本误拦截率不高于 1%。",
        "快速分类模型准确率不低于 98%、P99 不高于 200ms；主链路 P99 不高于 300ms（不含大模型推理）。",
        "单节点不低于 20 QPS，集群不低于 200 QPS，持续 30 分钟错误率不高于 0.1%。",
        "支持 3,000 并发会话和 300 管理控制台在线用户；CPU 峰值不高于 80%、内存峰值不高于 75%。",
    ]
    for item in targets:
        add_list_item(doc, item, bullet_id)


def add_scenario_value(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("4. 不同场景的产品价值", level=1)
    intro = doc.add_paragraph(
        "同一套安全能力在不同业务中承担的角色不同：有的强调在线低延迟阻断，有的强调异步深度分析，有的强调"
        "可追溯与审批。以下矩阵把升级能力转换为产品、业务和交付价值。"
    )
    intro.paragraph_format.space_after = Pt(8)
    add_table(
        doc,
        ["场景", "典型风险", "升级能力", "带来的价值"],
        [
            ["AI API / 模型网关", "接入协议不一、越权、流式内容绕检、突发流量", "统一 REST/SSE/WS/gRPC、签名上下文、配额并发、流式提交门", "业务无需重复建设护栏；可按应用切换策略，降低模型接入和审计成本"],
            ["保险客服与销售助手", "提示注入、收益夸大、虚假承诺、隐私泄露、错误拒答", "推理攻击、保险合规基线、DLP、Judge、代答/复核动作", "降低销售误导和客户信息泄露风险，同时把高风险回答稳定引导至人工"],
            ["保单/理赔材料审核", "文档、截图、扫描件中的隐私、隐藏指令或伪装内容", "分片上传、多视图 OCR/VLM、文档页级证据、图文联合决策", "让非结构化材料进入同一治理流程，提升证据定位、复核和批量处理能力"],
            ["录音与视频质检", "倍速、低音量、噪声、反向语音、闪帧和画音组合攻击", "多视图 ASR、声学异常、视频抽帧、OCR/视觉、时序跨模态融合", "从“只看转写文本”升级为画面、字幕、音轨的联合安全判断，支持时间戳回放"],
            ["企业知识库 / RAG", "恶意文档注入、来源不可追溯、越权检索、低信任内容改写系统规则", "入库扫描、来源版本、ACL、污染标记、信任边界与检索审计", "减少知识污染和间接注入；问题可回溯到原文档与片段，便于纠错和问责"],
            ["Agent 与自动化工具", "越权调用、参数注入、跨租户访问、数据外发或高风险动作", "工具白名单、Schema、资源范围、审批令牌、返回结果复检", "把“模型建议”与“真实执行”隔离，高风险动作可控、可审、可追踪"],
            ["安全运营与合规审计", "告警碎片化、证据不足、日志可篡改、无法重建全链路", "检测会话/发现项/事件事务、SLA、审计链、可信时间、Trace 与报表", "从单点命中升级为事件闭环，提升研判效率并为内审、等保和取证提供基础"],
            ["私有化与信创交付", "离线环境、组件异构、镜像来源、跨域网络和硬件差异", "Helm/Kustomize、mTLS、镜像摘要、SBOM、离线包、多架构与调度抽象", "降低交付重复劳动，提升环境一致性和供应链可追溯性，为国产化适配保留接口"],
        ],
        [1350, 2240, 2900, 2870],
        font_size=8.45,
    )

    doc.add_heading("场景化组合建议", level=2)
    combinations = [
        ("在线低延迟", "网关 + 快速规则/模型 + 流式门 + DLP；大文件和深度多模态转异步。"),
        ("高风险保险问答", "输入检测 + 推理攻击 + 保险规则 + 依据核验 + Judge + 人工复核。"),
        ("异步材料审查", "对象分片 + 内容校验 + OCR/VLM/ASR + 跨模态融合 + 事件/证据导出。"),
        ("Agent 执行", "RAG 信任边界 + 工具许可 + 高风险审批 + 结果复检 + 全链路审计。"),
    ]
    for label, value in combinations:
        add_label_paragraph(doc, label, value, after=5)


def add_verification_and_boundaries(doc: Document, bullet_id: int, number_id: int) -> None:
    add_page_break(doc)
    doc.add_heading("5. 验证证据与质量结论", level=1)
    add_callout(
        doc,
        "验证结论",
        "本次本地自动化、构建、数据库集成与前后端人工流程模拟均已通过；它们证明代码质量与联通性，"
        "但不能替代目标模型效果、真实流量性能、生产高可用和客户签名验收。",
        fill=PALE_TEAL,
        accent=TEAL,
    )
    add_table(
        doc,
        ["证据类别", "当前结果", "证据或命令"],
        [
            ["契约与兼容", "通过", "pnpm contracts:check"],
            ["TypeScript / ESLint", "通过", "pnpm validate / pnpm lint"],
            ["单元与覆盖率", "76 文件、300 用例通过", "pnpm test:coverage"],
            ["生产构建", "Next.js 与 Media Analyzer 均通过", "pnpm build / pnpm analyzer:build"],
            ["数据库集成", "PASS：29 迁移、12 关系、租户/谱系/审计通过", "pnpm test:integration"],
            ["前后端联通", "Playwright 10/10，通过桌面与移动端人工流程模拟", "pnpm test:e2e"],
            ["SDK 基线", "Python 3 项通过；Java/Go 在平台基线验证通过", "test:sdk-python / Maven / Go 基线报告"],
            ["需求追溯", "101/101 路径有效；0 项正式签名验收", "acceptance/generated/coverage-audit.json"],
        ],
        [2040, 3100, 4220],
        font_size=9,
        status_col=1,
    )

    doc.add_heading("覆盖率解读", level=2)
    for item in [
        "整体行覆盖率 27.72% 说明仓库仍有大量 UI、路由编排、数据库服务和真实子进程链未被单元测试直接执行。",
        "本次新增的推理攻击检测、时序融合、跨模态融合和策略加载具有 95%-100% 的行覆盖率，核心决策逻辑风险得到重点控制。",
        "Media Analyzer 的 audio-video.ts 行覆盖率为 14.28%，原因是 FFmpeg、OCR/VLM/ASR 真实命令集成不适合用纯单元测试替代；应以容器化 POC 补齐。",
        "前后端 E2E 与数据库集成覆盖了单元测试缺失的主要用户路径和持久化约束，但仍需继续提高路由与 UI 自动化覆盖。",
    ]:
        add_list_item(doc, item, bullet_id)

    doc.add_heading("6. 已知边界与未关闭事项", level=1)
    add_table(
        doc,
        ["边界", "当前状态", "关闭条件"],
        [
            ["多模态模型效果", "适配器、视图、时间轴和融合完成；目标模型未签名验收", "接入客户 OCR/VLM/ASR/音频模型，执行扰动、画音组合和格式专项 POC"],
            ["检测效果指标", "2,000 黑/白样本评分框架完成，无客户数据结果", "冻结数据集哈希与标注，执行盲测并提交签名报告"],
            ["性能与长稳", "k6 脚本和阈值已定义，无目标全策略压测报告", "在生产同构环境完成 P99、QPS、并发、资源和 72 小时长稳"],
            ["企业安全平台", "接口与适配边界具备，IdP/KMS/TSA/SOC 等未联调", "取得地址、证书、凭据与责任人，完成失败补偿和权限隔离测试"],
            ["高可用与灾备", "Helm/运行手册/拓扑模板具备，未做目标环境故障注入", "执行节点/依赖故障、恢复演练与 RPO/RTO 证据签名"],
            ["信创互认证", "多架构和调度抽象具备，昇腾覆盖文件为模板", "在指定 CPU/NPU/OS/数据库组合完成安装、性能和互换 POC"],
            ["正式验收", "101 项均待签名证据", "按需求逐项归档原始日志、Trace、配置摘要和客户签名"],
        ],
        [2040, 3520, 3800],
        font_size=8.8,
        status_col=1,
    )

    doc.add_heading("7. 建议上线与价值验证路径", level=1)
    steps = [
        "内部灰度：选择 1-2 个低风险应用，启用只审计/告警模式，建立误报、漏报、人工复核和成本基线。",
        "目标联调：接入企业 IdP/KMS/TSA、真实 OCR/VLM/ASR、模型网关、RAG 与 Agent，关闭接口和数据边界缺口。",
        "效果与性能验收：冻结黑白样本、策略和模型版本，执行检测效果、P99、QPS、并发、资源、长稳和故障注入。",
        "分级生产：先启用高置信度阻断，再按业务线扩大策略；对保险高风险回答和工具执行保持人工复核与审批。",
    ]
    for step in steps:
        add_list_item(doc, step, number_id)

    doc.add_heading("建议持续跟踪的产品指标", level=2)
    for metric in [
        "安全效果：攻击拦截率、召回率、精准率、白样本误拦率、跨模态协同攻击检出率。",
        "用户体验：P50/P95/P99 附加时延、人工复核等待时间、拒答率、代答采用率。",
        "运营效率：事件自动归并率、SLA 达成率、平均研判时间、误报反馈到策略发布的周期。",
        "资源成本：每千次文本检测成本、每分钟音视频分析成本、队列等待、模型并发和批处理利用率。",
        "治理成熟度：策略审批/回滚成功率、审计链验证率、删除证明覆盖率、需求签名验收完成率。",
    ]:
        add_list_item(doc, metric, bullet_id)


def add_appendices(doc: Document) -> None:
    doc.add_heading("附录 A  升级功能总清单", level=1)
    rows = [
        ["F-01", "接入", "REST / SSE / WebSocket / gRPC 统一网关", "已完成，待目标网络 POC"],
        ["F-02", "接入", "签名请求上下文、配额与分布式并发租约", "已完成，待 Redis/身份 POC"],
        ["F-03", "契约", "Guard v1 五种生成物与向后兼容门禁", "已完成并通过"],
        ["F-04", "SDK", "Java / Python / Go / TypeScript SDK", "已完成，待客户集成 POC"],
        ["F-05", "检测", "有界多视图归一化与提示攻击检测", "已完成并覆盖"],
        ["F-06", "检测", "分步推理攻击与推理过程提取", "已完成并覆盖"],
        ["F-07", "检测", "多轮累计风险、步骤证据和阻断点", "已完成并覆盖"],
        ["F-08", "行业", "结构化 DLP、资源滥用与保险合规基线", "基线完成，需数据校准"],
        ["F-09", "图片/文档", "多视图 OCR/VLM、图文联合决策", "链路完成，需目标模型 POC"],
        ["F-10", "音频", "多视图 ASR、声学异常、时间戳回映", "链路完成，需目标模型 POC"],
        ["F-11", "视频", "抽帧、OCR/视觉、音轨与时间轴融合", "链路完成，需目标模型 POC"],
        ["F-12", "多模态", "批次、抽帧、置信度、复核/阻断与时间窗策略", "已完成并覆盖"],
        ["F-13", "RAG", "来源、版本、ACL、污染标记和检索审计", "已完成，待真实 RAG POC"],
        ["F-14", "Agent", "工具白名单、Schema、资源范围、审批与结果复检", "已完成，待真实 Agent POC"],
        ["F-15", "策略", "签名策略包、审批分离、影子/灰度/回滚", "核心完成，部分产品功能待完善"],
        ["F-16", "评测", "异步评测、指标、重试与 badcase", "核心完成，标准报告待完善"],
        ["F-17", "运营", "事件 SLA、责任人、状态机与时间线", "已完成，聚合/反馈闭环待完善"],
        ["F-18", "审计", "哈希链、追加写、可信时间与独立验证", "已完成，待企业 TSA/SOC POC"],
        ["F-19", "数据", "租户隔离、谱系、保留与删除证明", "核心完成，外部存储闭环待完善"],
        ["F-20", "标识", "文本/图片/音频/视频显式与隐式内容标识", "已完成，待媒体 POC"],
        ["F-21", "部署", "Helm/Kustomize、mTLS、NetworkPolicy、PDB/HPA", "模板完成，待目标 K8s POC"],
        ["F-22", "供应链", "离线包、镜像摘要、Cosign、SBOM、多架构", "流程完成，待交付环境验证"],
        ["F-23", "可观测", "指标、日志、Trace、Grafana 与告警规则", "已完成，待环境 POC"],
        ["F-24", "控制台", "策略、评测、事件、文档、历史与模型管理页面", "主要路径 E2E 通过"],
    ]
    add_table(
        doc,
        ["编号", "能力域", "升级项", "状态"],
        rows,
        [900, 1200, 4360, 2900],
        font_size=8.6,
        status_col=3,
    )

    doc.add_heading("附录 B  编制依据", level=1)
    sources = [
        "Git 版本范围：af45893..41e3729；平台升级提交 6790142；检测增强提交 41e3729。",
        "机器验收口径：acceptance/generated/coverage-audit.json、acceptance/poc-manifest.json、acceptance/performance-requirements.json。",
        "研发状态报告：输出/代码分析与升级/10_研发执行与验收状态报告_V1.1.md。",
        "关键实现：src/lib/guard-engine-v2、src/lib/multimodal、src/lib/media、services/media-analyzer、services/guard-gateway。",
        "生产交付：deploy/、docs/runbooks/、packages/contracts/、packages/sdk-*。",
        "本轮复测：pnpm test:coverage、pnpm analyzer:build、pnpm test:integration；生产构建与 E2E 采用同一代码基线的本轮结果。",
    ]
    for source in sources:
        p = doc.add_paragraph(style="Code Reference")
        p.add_run(source)

    add_callout(
        doc,
        "最终说明",
        "本报告总结的是当前代码基线已形成的产品能力与工程证据。任何效果、性能、可用性、安全测评或信创互认证结论，"
        "只有在目标环境取得不可变报告并由客户签名后，才应升级为“正式验收通过”。",
        fill=PALE_GOLD,
        accent=GOLD,
    )


def audit_document(doc: Document) -> None:
    section = doc.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.left_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.top_margin == Inches(1)
    assert section.bottom_margin == Inches(1)

    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        layout = tbl_pr.find(qn("w:tblLayout"))
        assert tbl_w is not None and tbl_w.get(qn("w:w")) == str(CONTENT_WIDTH_DXA)
        assert tbl_ind is not None and tbl_ind.get(qn("w:w")) == str(TABLE_INDENT_DXA)
        assert layout is not None and layout.get(qn("w:type")) == "fixed"
        grid_widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        assert sum(grid_widths) == CONTENT_WIDTH_DXA
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                assert tc_w is not None and int(tc_w.get(qn("w:w"))) == grid_widths[index]

    full_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    forbidden = ["TODO", "TBD", "待补充", "{{", "}}", ":codex-file-citation"]
    for token in forbidden:
        assert token not in full_text, f"Placeholder or internal token found: {token}"


def build_document(output: Path) -> None:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)

    doc.core_properties.title = "国舜大模型安全护栏大版本升级总结"
    doc.core_properties.subject = "GuardLLM 功能、性能与场景价值说明"
    doc.core_properties.author = "国舜大模型安全护栏项目组"
    doc.core_properties.keywords = "GuardLLM, 大模型安全, 多模态, 推理攻击, 企业级护栏"
    doc.core_properties.comments = "基于 main@41e3729 编制"
    doc.core_properties.created = datetime(2026, 9, 2, 21, 30, 0)
    doc.core_properties.modified = datetime(2026, 9, 2, 21, 30, 0)

    bullet_id = add_numbering_definition(doc, kind="bullet", abstract_id=80, num_id=80)
    number_id = add_numbering_definition(doc, kind="decimal", abstract_id=81, num_id=81)

    with tempfile.TemporaryDirectory(prefix="guardllm-upgrade-doc-") as temp_dir:
        architecture_path = Path(temp_dir) / "architecture.png"
        create_architecture_image(architecture_path)
        add_cover(doc)
        add_executive_summary(doc, bullet_id)
        add_scope_and_architecture(doc, bullet_id, architecture_path)
        add_feature_sections(doc, bullet_id)
        add_performance_section(doc, bullet_id)
        add_scenario_value(doc)
        add_verification_and_boundaries(doc, bullet_id, number_id)
        add_appendices(doc)

        audit_document(doc)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_document(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
