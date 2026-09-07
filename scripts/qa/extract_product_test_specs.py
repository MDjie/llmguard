from __future__ import annotations

import argparse
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def iter_blocks(parent: DocumentType) -> Iterator[Paragraph | Table]:
    for child in parent.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def paragraph_record(paragraph: Paragraph) -> dict[str, Any]:
    properties = paragraph._p.pPr
    numbering = None
    if properties is not None and properties.numPr is not None:
        num_id = properties.numPr.numId
        level = properties.numPr.ilvl
        numbering = {
            "numId": num_id.val if num_id is not None else None,
            "level": level.val if level is not None else None,
        }
    return {
        "type": "paragraph",
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "numbering": numbering,
    }


def table_record(table: Table, index: int) -> dict[str, Any]:
    return {
        "type": "table",
        "index": index,
        "rows": [[cell.text for cell in row.cells] for row in table.rows],
    }


def extract_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    blocks: list[dict[str, Any]] = []
    table_index = 0
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            blocks.append(paragraph_record(block))
        else:
            table_index += 1
            blocks.append(table_record(block, table_index))

    headers: list[list[str]] = []
    footers: list[list[str]] = []
    for section in document.sections:
        headers.append([p.text for p in section.header.paragraphs if p.text.strip()])
        footers.append([p.text for p in section.footer.paragraphs if p.text.strip()])

    return {
        "path": str(path),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "section_count": len(document.sections),
        "headers": headers,
        "footers": footers,
        "blocks": blocks,
    }


def extract_xlsx(path: Path) -> dict[str, Any]:
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    sheets: list[dict[str, Any]] = []
    try:
        for formula_sheet in formulas.worksheets:
            value_sheet = values[formula_sheet.title]
            cells: list[dict[str, Any]] = []
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.value is None and not cell.comment and not cell.hyperlink:
                        continue
                    cached = value_sheet[cell.coordinate].value
                    cells.append(
                        {
                            "coordinate": cell.coordinate,
                            "value": json_value(cell.value),
                            "cached_value": json_value(cached),
                            "data_type": cell.data_type,
                            "number_format": cell.number_format,
                            "comment": cell.comment.text if cell.comment else None,
                            "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                            "hidden_row": bool(formula_sheet.row_dimensions[cell.row].hidden),
                            "hidden_column": bool(formula_sheet.column_dimensions[cell.column_letter].hidden),
                        }
                    )
            sheets.append(
                {
                    "name": formula_sheet.title,
                    "state": formula_sheet.sheet_state,
                    "max_row": formula_sheet.max_row,
                    "max_column": formula_sheet.max_column,
                    "merged_ranges": [str(rng) for rng in formula_sheet.merged_cells.ranges],
                    "freeze_panes": str(formula_sheet.freeze_panes) if formula_sheet.freeze_panes else None,
                    "auto_filter": formula_sheet.auto_filter.ref,
                    "cells": cells,
                }
            )
    finally:
        formulas.close()
        values.close()
    return {"path": str(path), "sheets": sheets}


def write_docx_markdown(data: dict[str, Any], path: Path) -> None:
    lines = ["# DOCX extracted content", ""]
    for block in data["blocks"]:
        if block["type"] == "paragraph":
            text = block["text"].strip()
            if text:
                lines.append(f"[{block['style']}] {text}")
        else:
            lines.extend(["", f"## Table {block['index']}", ""])
            for row in block["rows"]:
                lines.append(" | ".join(value.replace("\n", " / ") for value in row))
    path.write_text("\n".join(lines), encoding="utf-8")


def write_xlsx_markdown(data: dict[str, Any], path: Path) -> None:
    lines = ["# XLSX extracted content", ""]
    for sheet in data["sheets"]:
        lines.extend(
            [
                f"## {sheet['name']}",
                "",
                f"State: {sheet['state']}; Used range: {sheet['max_row']} rows x {sheet['max_column']} columns",
                "",
            ]
        )
        by_row: dict[int, list[tuple[int, str, Any]]] = {}
        for cell in sheet["cells"]:
            coordinate = cell["coordinate"]
            row_number = int("".join(ch for ch in coordinate if ch.isdigit()))
            column_letters = "".join(ch for ch in coordinate if ch.isalpha())
            column_number = 0
            for ch in column_letters:
                column_number = column_number * 26 + ord(ch.upper()) - 64
            by_row.setdefault(row_number, []).append((column_number, coordinate, cell["value"]))
        for row_number in sorted(by_row):
            parts = []
            for _, coordinate, value in sorted(by_row[row_number]):
                rendered = "" if value is None else str(value).replace("\n", " / ")
                parts.append(f"{coordinate}={rendered}")
            lines.append(" | ".join(parts))
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    docx_data = extract_docx(args.docx)
    xlsx_data = extract_xlsx(args.xlsx)
    inventory = {"docx": docx_data, "xlsx": xlsx_data}

    (args.output_dir / "source-inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_docx_markdown(docx_data, args.output_dir / "manual-content.md")
    write_xlsx_markdown(xlsx_data, args.output_dir / "test-cases-content.md")

    print(
        json.dumps(
            {
                "docx": {
                    "paragraphs": docx_data["paragraph_count"],
                    "tables": docx_data["table_count"],
                    "inline_shapes": docx_data["inline_shape_count"],
                    "sections": docx_data["section_count"],
                },
                "xlsx": [
                    {
                        "name": sheet["name"],
                        "state": sheet["state"],
                        "rows": sheet["max_row"],
                        "columns": sheet["max_column"],
                        "nonempty_cells": len(sheet["cells"]),
                    }
                    for sheet in xlsx_data["sheets"]
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
