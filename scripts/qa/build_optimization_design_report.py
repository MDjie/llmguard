from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("输出/产品资料/2026-09-04/03_GuardLLM_安全检测能力优化设计报告_V1.0_2026-09-04.docx")

NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Inches(widths[index])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[column_index])
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) <= 18 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            for run in paragraph.runs:
                run.font.size = Pt(8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.add_run(item)
        paragraph.paragraph_format.space_after = Pt(3)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)
        paragraph.paragraph_format.space_after = Pt(3)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    set_keep_with_next(paragraph)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(14)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.2)


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.text = "GUARDLLM  安全检测能力优化设计报告  V1.0"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(8)
            run.font.color.rgb = BLACK

        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("内部评审和实施设计  |  2026-09-04  |  第 ")
        run.font.size = Pt(8)
        page_run = paragraph.add_run()
        page_run.font.size = Pt(8)
        add_field(page_run, " PAGE ")
        end = paragraph.add_run(" 页")
        end.font.size = Pt(8)


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    configure_styles(doc)
    add_header_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("GuardLLM 安全检测能力优化设计报告")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("输入防护 输出管控 敏感词治理与文本攻击检测增强").bold = True
    subtitle.runs[0].font.size = Pt(14)
    doc.add_paragraph()
    add_table(doc, ["项目", "内容"], [
        ["文档版本", "V1.0"],
        ["编制日期", "2026-09-04"],
        ["设计基线", "main@69639dc 与 2026-09-04 全量自动化测试报告"],
        ["适用范围", "GuardLLM Version A 软件网关及后续兼容扩展"],
        ["目标读者", "产品负责人 技术负责人 安全算法 后端平台 测试 运维与交付团队"],
        ["文档状态", "详细设计和实施评审稿"],
    ], [1.45, 5.4])
    add_body(doc, "本报告给出可直接进入研发排期的优化设计。首要任务是恢复签名策略包的可用发布链路；在此基础上，将现有规则、规范化、语义分类、多模态融合、会话安全记忆和资源控制能力整合为统一的双向安全决策流水线。设计重点增强角色扮演、目标劫持、提示词泄露、编码与字符混淆、图文协同攻击、多轮递进诱导、算力耗尽，以及输出侧政治、色情和保险行业合规、隐私脱敏与分级代答能力。")
    doc.add_page_break()

    add_heading(doc, "文档目录", 1)
    add_numbered(doc, [
        "执行摘要",
        "测试基线与问题诊断",
        "优化目标与设计原则",
        "目标安全架构",
        "输入侧防护增强设计",
        "输出侧管控增强设计",
        "敏感词与规则治理平台",
        "策略包和部署可用性整改",
        "接口 数据与审计设计",
        "性能 可靠性与安全设计",
        "测试验证与验收标准",
        "实施计划与工作包",
        "风险 回滚与运营机制",
        "结论",
    ])

    add_heading(doc, "1 执行摘要", 1)
    add_body(doc, "测试表明，当前代码质量基线良好：495 个 Vitest 用例、数据库集成、10 个浏览器 E2E、SDK、Java Gateway 和生产构建均通过。但部署实例没有可执行的签名策略包，32 个 16 维运行时样本全部返回 POLICY_LOAD_FAILED；158 条正式用例中仅 13 条形成完整本地闭环，32 条失败，102 条仅部分覆盖，11 条受外部条件阻塞。当前版本不具备发布或客户验收条件。")
    add_body(doc, "优化工作必须分两层推进。第一层是可用性底座：完成签名密钥隔离、默认策略包编译审批激活、应用绑定、策略就绪探针以及只读容器缓存整改。第二层是能力增强：建设可版本化敏感词和攻击规则中心，以规范化视图、快速词法匹配、上下文分类、语义分类、多模态融合和 Judge 复核形成分层检测，并在输入和输出方向执行不同的动作策略。")
    add_table(doc, ["优先级", "目标", "完成标志"], [
        ["P0", "恢复 Guard 运行时与部署就绪", "默认签名策略包可验证可激活 32 个维度样本均进入判定 App 不再出现缓存异常"],
        ["P1", "增强文本攻击和敏感词检测", "注入 混淆 多语种 结构化 DLP 与保险合规形成稳定离线回归"],
        ["P2", "完成多模态 多轮和资源攻击闭环", "图文 音视频 会话风险与生命周期预算统一决策"],
        ["P3", "完成目标环境验收", "独立盲测 性能 第三方联调 国产化和可靠性报告具名签收"],
    ], [0.7, 2.0, 4.15])

    add_heading(doc, "2 测试基线与问题诊断", 1)
    add_heading(doc, "2.1 测试结果概览", 2)
    add_table(doc, ["测试域", "结果", "设计含义"], [
        ["代码和构建", "通过", "可在现有 TypeScript Next.js 架构上增量演进 不需要重写平台"],
        ["数据库集成", "通过", "现有迁移 租户隔离 数据血缘和只追加审计可继续复用"],
        ["浏览器 E2E", "通过", "登录 鉴权 导航 注销和健康检查可作为回归骨架"],
        ["16 维运行时", "32 条全部失败", "策略发布可用性是所有检测能力的前置阻断"],
        ["验收追踪", "101 条待签收", "代码存在不能替代目标环境和外部证据"],
        ["兼容性", "15 个目标待 POC", "设计必须保留适配器和证据冻结机制"],
    ], [1.55, 1.35, 3.95])

    add_heading(doc, "2.2 根因与结构性缺口", 2)
    add_bullets(doc, [
        "策略运行时断点。数据库存在启用的策略 Profile，但 policy_bundles 为 0；应用容器缺少策略签名公钥、私钥和 Key ID，Compose 公共运行环境没有映射相关变量。",
        "能力碎片化。规范化、内置检测器、语义分类、PII 脱敏、安全改写、多模态融合和资源预算均已存在，但没有在已激活签名策略包中形成可部署、可验证、可回滚的完整组合。",
        "敏感词治理不足。现有 keyword_categories 和 keyword_rules 可支持基础管理，但缺少词典版本、变体生成、上下文白名单、灰度发布、效果指标和 badcase 闭环。",
        "输出动作分散。API 路由内存在脱敏和改写逻辑，guardrail 目录也有独立实现，动作语义、模板版本、证据留存和复检行为需要统一。",
        "验收证据不足。多模态、Agent、评测、扫描、性能和外部集成大多只有代码或框架证据，缺少目标环境闭环。",
    ])
    add_heading(doc, "2.3 可复用现有能力", 2)
    add_table(doc, ["现有模块", "可复用能力", "主要增强方向"], [
        ["guard-engine-v2 normalization", "NFKC 零宽字符 HTML URL Base64 Base32 Hex QP ROT13 和混淆字符视图", "增加切片重组 拼音及跨语种变体 编码图预算与可疑度评分"],
        ["builtin detectors", "注入 资源滥用 保险合规 结构化 DLP", "词典外置 规则版本化 多语言模板 上下文校准和覆盖度指标"],
        ["semantic classifier", "模型身份摘要 校准 阈值 SHADOW ENFORCE 和 fail closed", "多标签集成 对抗训练 分类别阈值 漂移监控"],
        ["session context", "安全记忆 多轮风险合并 版本冲突处理", "意图图 风险衰减 跨轮证据链和会话速率限制"],
        ["multimodal fusion", "用户文本 OCR 图像发现项和协同攻击判断", "ASR 视频时间轴 QR 条码 字幕和跨模态推理诱导"],
        ["PII masker rewrite engine", "结构化脱敏和安全响应基础", "统一变换引擎 行业数据类型 模板版本与二次复检"],
        ["resource control", "复杂度估算 生命周期预算 公平调度", "入站预估 动态限额 取消传播和租约恢复"],
    ], [1.6, 2.45, 2.8])

    add_heading(doc, "3 优化目标与设计原则", 1)
    add_heading(doc, "3.1 目标能力", 2)
    add_bullets(doc, [
        "输入侧对直接和间接提示注入、角色扮演、权限抬升、目标劫持、提示词及推理过程泄露进行语义识别。",
        "对编码、切片、同形字符、零宽字符、拼音、谐音、多语种翻译、标点插入和跨载体拆分等绕过进行深度还原。",
        "对图片、OCR、二维码、音频、视频字幕和帧级视觉信号执行单模态检测及图文协同推理。",
        "对超长输出、递归、组合爆炸、工具扇出、长上下文和多轮递进攻击实施预算控制和会话级拦截。",
        "输出侧覆盖政治、色情低俗及保险行业特有合规风险，并自动保护客户隐私、健康、金融和内部敏感数据。",
        "以 ALLOW WARN MASK REWRITE REQUIRE_REVIEW SAFE_RESPONSE BLOCK 七级动作实现红线标准拒答与一般风险自定义代答。",
    ])
    add_heading(doc, "3.2 设计原则", 2)
    add_table(doc, ["原则", "要求"], [
        ["策略先于代码", "检测词典 阈值 模型摘要 动作模板和失败策略全部进入签名策略包 不依赖路由硬编码"],
        ["分层检测", "低成本规则优先 语义和 Judge 按风险升级 多模态与会话上下文按需触发"],
        ["最小化证据", "默认只保存偏移 掩码预览 HMAC 模型和策略版本 不保存原始敏感文本"],
        ["可解释可复核", "每个动作返回风险类型 来源模态 原始偏移 规则或模型版本和决策路径"],
        ["失败安全", "关键检测器 策略验证和输出复检失败时不得静默放行"],
        ["租户隔离", "词典 模板 阈值 数据集和反馈按租户与应用隔离 全部经过作用域校验"],
        ["灰度可回滚", "新规则和模型先 SHADOW 再 CANARY 后 ACTIVE 保留最近可用策略包"],
    ], [1.5, 5.35])

    add_heading(doc, "4 目标安全架构", 1)
    add_heading(doc, "4.1 双向决策流水线", 2)
    add_table(doc, ["阶段", "输入侧处理", "输出侧处理"], [
        ["接入与预算", "认证 租户应用作用域 请求大小 Token 和复杂度预算", "输出 Token 流式窗口和最大生成预算"],
        ["内容展开", "文本 附件 OCR ASR 视频帧 QR 条码 RAG 和工具上下文", "模型文本 图片 音频 视频 工具结果和引用来源"],
        ["规范化", "Unicode 编码还原 切片重组 多语种和字符骨架", "同样执行规范化 并保留原始偏移和模态来源"],
        ["快速检测", "敏感词 Aho Corasick 安全正则 校验和及结构化实体", "政治 色情 保险话术 PII 凭证和内部数据"],
        ["高级检测", "语义注入 多轮攻击 跨模态协同和资源攻击", "语义合规 DLP 上下文泄露 幻觉式承诺和引用完整性"],
        ["聚合与复核", "规则 语义 会话 模态风险融合 必要时 Judge", "按行业 方向和对象执行更严格输出阈值"],
        ["动作执行", "允许 警告 阻断或要求人工复核", "脱敏 改写 自定义代答 标准拒答 阻断并二次复检"],
        ["证据与运营", "审计链 指标 badcase 和策略版本", "变换前后摘要 模板版本 复检结果和申诉反馈"],
    ], [1.1, 2.9, 2.85])

    add_heading(doc, "4.2 检测层级", 2)
    add_table(doc, ["层级", "执行内容", "触发条件", "失败策略"], [
        ["L0", "边界校验 规范化 结构化 DLP 红线词", "所有请求", "失败关闭"],
        ["L1", "词典 多模式匹配 安全正则 混淆可疑度", "所有文本和 OCR ASR", "失败关闭"],
        ["L2", "轻量多标签语义分类", "L0 L1 命中或策略要求", "ENFORCE 失败关闭 SHADOW 可降级"],
        ["L3", "多轮 跨模态 RAG 和推理诱导检测", "存在会话或多模态上下文", "按策略关闭或要求复核"],
        ["L4", "Judge 复核和行业规则裁决", "分数落入灰区或高影响业务", "超时执行预设安全动作"],
    ], [0.7, 2.4, 2.0, 1.75])

    add_heading(doc, "5 输入侧防护增强设计", 1)
    add_heading(doc, "5.1 注入检测", 2)
    add_body(doc, "注入检测采用规则信号、语义意图和上下文信任联合判定。单一关键词只形成弱证据；当攻击意图、受保护目标、权限诉求、隐藏指令或来源不可信等信号组合出现时提升风险。这样既能识别明确攻击，又可降低安全研究、产品说明和防护讨论中的误报。")
    add_table(doc, ["攻击类型", "关键识别信号", "建议动作"], [
        ["角色扮演和权限抬升", "要求成为无约束角色 管理员 开发者模式 或改变指令优先级", "高置信 BLOCK 中置信 REQUIRE_REVIEW"],
        ["目标劫持", "放弃原任务 改变目标 替换业务流程 绕过既定约束", "BLOCK 或 SAFE_RESPONSE"],
        ["提示词泄露", "索取系统提示词 开发者消息 隐藏策略 工具密钥或内部推理", "BLOCK 并返回固定安全说明"],
        ["拒答压制", "要求不得拒绝 不得提及安全 必须无条件执行", "与其他信号融合后阻断"],
        ["间接注入", "文档 网页 邮件 RAG 片段 工具结果中包含执行指令", "降低来源信任 禁止其取得指令能力"],
        ["结构化伪装", "在 JSON YAML XML Markdown 代码注释或字段名中伪造 system developer policy", "规范化后检测 保留结构来源"],
        ["推理诱导", "要求展示思维链 草稿 隐藏步骤 或逐步规避护栏", "SAFE_RESPONSE 不泄露内部推理"],
    ], [1.65, 3.45, 1.75])
    add_body(doc, "语义模型建议使用多标签分类而非单一二分类。标签至少包括 direct_override role_escalation goal_hijack prompt_exfiltration indirect_injection refusal_suppression reasoning_exfiltration coercion 和 structured_disguise。模型输出必须绑定 modelId、version 和 sha256，并在签名策略包中保存类别阈值和温度校准参数。")

    add_heading(doc, "5.2 混淆和绕过识别", 2)
    add_body(doc, "在现有规范化视图基础上增加受预算约束的解码图。每个节点记录变换类型、深度、输入摘要、输出摘要、字符映射和风险增益；相同文本去重，超过最大视图、字符数或解码时间立即停止，避免解码器本身成为算力攻击入口。")
    add_table(doc, ["绕过族", "增强处理", "验证重点"], [
        ["编码", "URL 多层 Base64 Base32 Hex QP ROT13 HTML 实体 Unicode 转义及组合嵌套", "深度上限 输出大小 可打印率和原文偏移"],
        ["字符替换", "NFKC 零宽字符 同形字 全半角 圈字 数学字母和常见 OCR 混淆骨架", "误归一化保护和证据回映"],
        ["切片", "按空格 标点 换行 Markdown 表格 代码块 多消息 多附件重组候选短语", "只在同一可信边界或会话窗口内重组"],
        ["多语种", "语言识别 中英及重点语种语义分类 翻译对齐 多语言攻击模板", "保持原语言证据 禁止翻译成为单点判定"],
        ["拼音谐音", "拼音首字母 全拼 数字谐音 近音字和常见网络变体", "限定敏感类别和上下文 防止高误报"],
        ["视觉文本", "OCR 多角度 多尺度 字符框排序 QR 条码和截图文字提取", "区域坐标 置信度和跨区域拼接"],
    ], [1.2, 3.7, 1.95])
    add_body(doc, "混淆可疑度作为独立特征进入聚合，而不是直接等价于恶意。建议特征包括规范化前后编辑距离、脚本混用比例、零宽字符密度、解码深度、异常分隔符密度和跨载体重组次数。只有可疑度与攻击语义或红线内容共同出现时才提升到阻断。")

    add_heading(doc, "5.3 多模态防护", 2)
    add_table(doc, ["模态", "处理链", "新增检测"], [
        ["图片", "文件校验 解码 OCR 视觉分类 QR 条码 区域融合", "图中文字注入 有害图像 图文组合后语义升级 隐写异常提示"],
        ["音频", "格式校验 ASR 说话人片段 时间戳 音频事件分类", "口播敏感内容 隐蔽指令 音频与文本任务冲突"],
        ["视频", "镜头切分 风险自适应采样 OCR ASR 帧分类 时间轴融合", "短帧闪现 字幕与画面组合 推理诱导和跨时段累积"],
        ["文档", "解压防护 转换 OCR 嵌入对象 宏和链接隔离", "扩展名欺骗 文档内间接注入 隐藏层和页间切片"],
    ], [1.0, 3.25, 2.6])
    add_body(doc, "融合层继续复用 user image combined 三路决策，但扩展为 text OCR visual ASR timeline 和 metadata 六类来源。若单个模态看似安全、组合后动作等级提升，则标记 cooperative_attack；证据必须定位到文件、页码、时间段、帧和区域。高风险文件在分析完成前不得进入模型上下文。")

    add_heading(doc, "5.4 新型威胁防护", 2)
    add_heading(doc, "5.4.1 算力耗尽", 3)
    add_bullets(doc, [
        "在读取完整正文前检查 Content Length、文件数量、页数、媒体时长、压缩比和递归归档深度。",
        "使用精确 Tokenizer 计算输入、历史、RAG、预期输出和检测器成本；Tokenizer 摘要必须与策略包一致。",
        "将 toolSteps、recursionDepth、browserTabs、processes、connections、files、ocrPages、mediaDurationSeconds 和 guardInferenceTokens 纳入生命周期预算。",
        "按租户、应用、主体、凭证和接口设置并发租约与公平队列；取消信号必须传播到 OCR、语义模型、Judge 和工具调用。",
        "超限返回可审计的 RESOURCE_BUDGET_EXCEEDED，不进入昂贵检测；关键护栏不可因资源不足而默认放行。",
    ])
    add_heading(doc, "5.4.2 多轮递进诱导", 3)
    add_body(doc, "在现有安全记忆上建立会话意图图。每轮记录目标、方法、受保护对象、隐藏意图、风险类别和动作，按时间和轮次衰减；单轮弱信号在跨轮形成方法加目标加规避链条时升级。安全记忆只保存摘要、标签和 HMAC，不保存不必要的原始隐私。")
    add_table(doc, ["状态", "条件", "动作"], [
        ["NORMAL", "无风险或低置信单点信号", "正常检测"],
        ["WATCH", "两轮出现相关弱信号 或存在规避和受保护目标", "提高采样和语义检测强度"],
        ["ESCALATED", "形成连续攻击链 或重复触发高风险", "收紧阈值 REQUIRE_REVIEW 或 SAFE_RESPONSE"],
        ["LOCKED", "明确越权 数据外泄或持续绕过", "会话级 BLOCK 冷却期和事件处置"],
    ], [1.0, 3.75, 2.1])

    add_heading(doc, "6 输出侧管控增强设计", 1)
    add_heading(doc, "6.1 全模态合规审核", 2)
    add_body(doc, "输出审核与输入检测共用基础引擎，但使用独立的 output policy。输出侧更关注生成结果是否可以向用户呈现，阈值、动作和模板不能简单复用输入侧。文本、图片、音频、视频和工具结果在提交给用户前都必须完成输出门控；流式输出采用缓冲窗口和提交门，风险片段不得先发送后撤回。")
    add_table(doc, ["合规域", "检测范围", "处置要求"], [
        ["政治合规", "按部署地区和客户规则配置政治人物 组织 事件 立场诱导 谣言和敏感表达类别", "分类分级 规则版本可追溯 高影响场景支持人工复核"],
        ["色情低俗", "露骨性内容 性暗示 低俗营销 未成年人相关内容 图像和视频成人内容", "未成年人或明确红线 BLOCK 一般低俗按行业策略代答"],
        ["保险误导", "保本保收益 保证理赔 隐瞒健康告知 弱化免责条款 高压营销 虚构监管背书", "标准纠偏答复 保留条款提示 禁止确定性承诺"],
        ["保险适当性", "未了解年龄 风险承受 健康和需求即给出确定产品建议", "要求补充信息或转人工 不输出个性化承诺"],
        ["违法和有害", "违法步骤 暴力仇恨 自伤 危险代码和欺诈", "按红线策略拒答并提供安全替代信息"],
    ], [1.25, 3.75, 1.85])
    add_body(doc, "政治合规词库必须由客户授权的合规责任人维护，并按司法辖区、行业和业务线版本化。系统提供类别、证据和动作框架，不将未经审批的临时词表固化到代码。")

    add_heading(doc, "6.2 隐私和内部敏感数据脱敏", 2)
    add_table(doc, ["数据类别", "识别方法", "默认变换"], [
        ["客户身份", "姓名 手机 邮箱 身份证 护照 地址 客户号及实体上下文", "保留必要前后缀或稳定令牌化"],
        ["保险业务", "保单号 理赔号 受益人 健康告知 病历 体检和核保结论", "字段级掩码 高敏健康数据默认全隐藏"],
        ["金融数据", "银行卡 账户余额 收入 征信 支付信息", "校验和识别后最小展示"],
        ["凭证秘密", "API Key Token Password Cookie 私钥 AK SK 和连接串", "全部替换并触发高风险事件"],
        ["内部数据", "内部定价 未公开产品 规则 模型提示词 架构和人员信息", "按数据目录标签执行掩码或阻断"],
    ], [1.2, 3.7, 1.95])
    add_body(doc, "统一脱敏引擎以实体跨度为输入，先处理重叠区间，再按策略执行 partial_mask full_mask tokenize redact 或 block。不得在 API 响应、日志、审计和异常堆栈中保存 original 字段；证据仅保留类别、偏移、掩码预览和内容 HMAC。变换完成后必须再次检测，确认敏感片段和上下文推断风险均已消除。")

    add_heading(doc, "6.3 智能干预和代答", 2)
    add_table(doc, ["风险等级", "动作", "用户体验"], [
        ["S0", "ALLOW", "正常返回"],
        ["S1", "WARN", "返回内容并显示合规提示 记录审计"],
        ["S2", "MASK", "掩码敏感字段 保持语义可用"],
        ["S3", "REWRITE", "删除违规承诺 补充必要限制条件和风险提示"],
        ["S4", "REQUIRE_REVIEW 或 SAFE_RESPONSE", "进入人工审核 或返回客户配置的一般风险代答"],
        ["S5", "BLOCK", "红线问题使用标准拒答 不回显风险原文"],
    ], [0.8, 2.1, 3.95])
    add_body(doc, "响应模板应进入独立的签名版本体系，包含 templateId、version、locale、industry、riskType、action、变量白名单和审批信息。红线模板由平台提供最小安全答复，一般风险模板允许租户按品牌和业务场景自定义。任何代答都要通过输出复检；若模板渲染失败或复检仍命中红线，则退回平台固定安全响应。")

    add_heading(doc, "7 敏感词与规则治理平台", 1)
    add_heading(doc, "7.1 词典分层", 2)
    add_table(doc, ["层级", "内容", "管理方式"], [
        ["平台红线库", "凭证泄露 明确违法 未成年人高风险和系统提示词泄露", "平台安全团队维护 租户不可降低强制动作"],
        ["行业库", "保险误导 健康告知 免责条款 适当性和监管表达", "行业专家和合规负责人双人审批"],
        ["租户库", "客户敏感人物 项目 产品 内部系统 数据字段和品牌规则", "租户隔离 灰度发布 可回滚"],
        ["应用库", "特定机器人 场景和渠道的词语 模板和例外", "应用负责人维护 安全管理员审批"],
        ["临时事件库", "突发事件和短期专项词表", "自动失效 必须设置有效期和责任人"],
    ], [1.25, 3.6, 2.05])

    add_heading(doc, "7.2 词条模型和匹配算法", 2)
    add_bullets(doc, [
        "词条字段包括 category、canonicalTerm、locale、severity、direction、industry、contexts、matchMode、mandatoryDeny、validFrom、validTo、owner 和 evidenceRequirement。",
        "自动生成大小写、全半角、简繁体、拼音、首字母、同形字、数字替换、空格标点切片和常见 OCR 变体；高误报变体必须人工批准。",
        "精确词和短语使用 Aho Corasick 自动机；结构化模式使用安全正则和校验和；近似匹配使用受长度限制的编辑距离或 token n gram；语义敏感主题交给分类器。",
        "每个命中携带 canonicalTermId、variantId、dictionaryReleaseId、viewId、原文偏移和 HMAC，支持解释、复盘和撤销。",
        "白名单不允许豁免 mandatoryDeny；普通例外必须限定租户、应用、方向、维度、有效期和审批人。",
    ])

    add_heading(doc, "7.3 上下文降误报", 2)
    add_body(doc, "敏感词不能脱离上下文直接阻断。引擎需要识别否定、引用、新闻报道、法律条文、安全研究、教育医疗、测试数据和用户授权范围。词法命中先产生 observation，再由上下文分类器判断 mention、endorsement、instruction、transaction、disclosure 等语义角色。只有风险角色与策略阈值匹配时执行高等级动作。")
    add_table(doc, ["上下文", "示例含义", "建议处理"], [
        ["提及或引用", "描述某敏感概念或引用合规条款", "记录或低风险 WARN"],
        ["倡导或承诺", "鼓励实施或作出不当保证", "REWRITE SAFE_RESPONSE 或 BLOCK"],
        ["操作指导", "给出可执行步骤和规避方法", "高风险 BLOCK"],
        ["数据披露", "输出可识别客户或内部秘密", "MASK 或 BLOCK"],
        ["防御研究", "检测 防御和合规培训语境", "降低误报但保持凭证和真实隐私保护"],
    ], [1.25, 3.65, 2.0])

    add_heading(doc, "7.4 发布与运营闭环", 2)
    add_numbered(doc, [
        "词条或规则提交后运行静态检查、ReDoS 检查、冲突检查、重复项检查和样本单测。",
        "生成不可变 dictionary release，并随策略包记录 sha256、提交人和审批人。",
        "先在 SHADOW 环境运行，比较命中率、误报率、动作差异和延迟。",
        "CANARY 按租户或流量比例放量，达到门槛后切换 ACTIVE；异常一键回滚到 last known good。",
        "将申诉、人工复核和 badcase 转为标注样本，进入周期性校准，不直接在线自学习。",
    ])

    add_heading(doc, "8 策略包和部署可用性整改", 1)
    add_heading(doc, "8.1 DEF 001 整改", 2)
    add_table(doc, ["整改项", "详细设计", "验收"], [
        ["密钥分权", "签名私钥只进入控制面签名任务或 KMS HSM 数据面和普通 Worker 仅持有公钥", "容器检查无私钥扩散 公钥摘要可审计"],
        ["本地引导", "增加 policy bootstrap 命令 生成或读取开发密钥 编译默认策略 运行测试 审批 激活并建立应用绑定", "空数据库部署后一次命令形成可用策略"],
        ["启动校验", "新增 /api/health/policy 校验默认 Profile 活动 Bundle 签名 模式和应用绑定", "无可用 Bundle 时 readiness 失败而不是健康假阳性"],
        ["原子发布", "编译 测试 审批 激活使用状态机和事务 更新 generation 后数据面原子切换", "并发请求只看到旧版或新版"],
        ["回滚", "缓存最近验证成功 Bundle 签名失败或数据库短时异常可按策略使用 LKG", "回滚事件和代际均可追踪"],
    ], [1.15, 4.15, 1.6])
    add_body(doc, "Compose 和生产部署清单必须显式映射 POLICY_SIGNING_PUBLIC_KEY 与 POLICY_SIGNING_KEY_ID；私钥不得放入公共 runtime environment。若当前单体 App 暂时承担签名职责，应以独立 Secret、最小权限路由和审计控制作为过渡，并在后续拆分为 signer job。")

    add_heading(doc, "8.2 DEF 002 整改", 2)
    add_body(doc, "只读 App 容器为 /app/.next/cache 配置容量受限的 tmpfs 或专用可写卷，并设置非 root 所有权、容量告警和清理策略。若生产不需要运行时图片优化，则关闭对应写盘路径。整改后日志不得再出现 mkdir ENOENT 或图片缓存 unhandled rejection。")

    add_heading(doc, "8.3 配置完整性门禁", 2)
    add_bullets(doc, [
        "构建阶段验证所有必需环境变量名称和 Secret 引用，不读取或打印 Secret 值。",
        "部署前验证策略公钥摘要、默认 Bundle ID、应用绑定、模型摘要、Tokenizer 摘要和词典发布摘要。",
        "健康检查拆分 live、db ready、policy ready、analyzer ready 和 worker lag，发布门禁要求全部 ready。",
        "策略发布、模型切换和词典更新均生成不可变证据清单，纳入审计链。",
    ])

    add_heading(doc, "9 接口 数据与审计设计", 1)
    add_heading(doc, "9.1 Guard 请求和响应", 2)
    add_table(doc, ["对象", "新增或统一字段", "目的"], [
        ["请求 context", "traceId requestId tenantId applicationId sessionId direction sourceType locale jurisdiction industry deadline policyBundleId tokenizerId", "固定安全上下文和策略选择"],
        ["内容 envelope", "envelopeId sourceId trustLevel instructionCapability parentIds modality artifactId page timeRange region contentHash", "建立来源信任和跨模态证据链"],
        ["响应 decision", "decisionId action riskLevel score confidence degraded policyPath latencyBreakdown", "统一动作和运行状态"],
        ["响应 finding", "riskType category severity evidenceRefs detector model rule dictionary versions", "可解释和可复核"],
        ["响应 transform", "type ranges templateId templateVersion outputHash recheckDecisionId", "追踪脱敏 改写和代答"],
    ], [1.2, 4.45, 1.25])

    add_heading(doc, "9.2 数据模型调整", 2)
    add_table(doc, ["对象", "调整建议"], [
        ["keyword_categories keyword_rules", "增加 releaseId locale direction industry severity contexts canonicalTerm variantType validFrom validTo owner mandatoryDeny"],
        ["dictionary_releases", "保存版本 sha256 状态 统计 提交 审批 激活和回滚信息"],
        ["response_templates", "保存风险类别 动作 语言 行业 变量白名单和审批状态"],
        ["detector_calibrations", "保存模型或规则版本 分类别阈值 温度 数据集摘要和效果指标"],
        ["badcase_feedback", "保存 decisionId 复核结论 原因 标签版本和处理状态 不复制原始敏感正文"],
        ["session_risk_state", "增加意图节点 累积分数 衰减参数 状态和最后证据序列"],
    ], [1.65, 5.2])

    add_heading(doc, "9.3 审计要求", 2)
    add_bullets(doc, [
        "记录策略包、词典、模型、Tokenizer、模板和检测器版本，不仅记录最终动作。",
        "记录输入与输出方向、来源模态、降级原因、超时、预算拒绝和人工复核结论。",
        "证据正文默认 HMAC 化和掩码，导出原始数据需双人审批并受保留期控制。",
        "审计链验证失败、策略签名失败或模板复检失败必须产生高优先级安全事件。",
    ])

    add_heading(doc, "10 性能 可靠性与安全设计", 1)
    add_heading(doc, "10.1 建议性能预算", 2)
    add_body(doc, "以下为待目标环境冻结的设计指标，不代表当前系统已经达到。测试必须记录 P50 P95 P99、QPS、并发、错误率和资源占用。")
    add_table(doc, ["路径", "建议 P95 增量", "控制手段"], [
        ["L0 规范化和结构化 DLP", "不高于 20 ms", "长度上限 视图去重 解码预算和预编译规则"],
        ["L1 敏感词与安全正则", "不高于 20 ms", "Aho Corasick 自动机 安全正则和缓存"],
        ["本地轻量语义分类", "不高于 150 ms", "批处理 模型常驻 并发限额和超时"],
        ["远程 Judge", "不高于 800 ms", "仅灰区触发 严格出站和响应大小限制"],
        ["文本主路径不含 Judge", "总增量不高于 200 ms", "并行 DAG 早停和分级执行"],
    ], [2.0, 1.5, 3.35])

    add_heading(doc, "10.2 可靠性", 2)
    add_bullets(doc, [
        "策略验证、关键红线检测器、输出提交门和脱敏复检均采用 fail closed。",
        "非关键辅助模型可按签名策略降级，但响应必须包含 degraded 和 reasonCodes。",
        "数据库不可用时只允许使用未过期的 last known good 策略；没有 LKG 时 readiness 失败。",
        "Worker 使用幂等任务、可见性超时、死信和重试上限；媒体和文档中间产物受保留期清理。",
        "缓存、模型、词典和策略均设置摘要校验，禁止未经验证的热更新。",
    ])

    add_heading(doc, "10.3 安全", 2)
    add_bullets(doc, [
        "签名私钥通过 KMS HSM 或只读 Secret 注入，签名操作具备最小权限和双人审批。",
        "语义模型、OCR、ASR 和 Judge 出站使用允许列表、TLS、请求大小限制和响应 Schema 校验。",
        "所有解码器、转换器和媒体分析进程在隔离容器中运行，限制 CPU 内存、文件、进程和网络。",
        "训练和评测数据删除真实标识符，使用独立密钥计算摘要，禁止将客户数据默认用于训练。",
    ])

    add_heading(doc, "11 测试验证与验收标准", 1)
    add_heading(doc, "11.1 自动化测试体系", 2)
    add_table(doc, ["层级", "新增测试", "门禁"], [
        ["单元", "规范化偏移 解码预算 词典自动机 实体重叠 动作模板和会话衰减", "每个规则和算法边界均有正反例"],
        ["变形测试", "同一攻击自动生成编码 切片 同形字 多语种和标点变体", "变形前后风险结论保持一致"],
        ["组件", "语义分类模型身份 阈值校准 多模态融合和失败策略", "响应 Schema 摘要和超时行为可复核"],
        ["集成", "策略编译 审批 激活 绑定 回滚 数据库和 Worker", "空库部署可自动形成可用默认策略"],
        ["E2E", "输入 输出 流式 RAG Agent 文档 图片 音频 视频和代答", "真实部署容器完成完整动作闭环"],
        ["盲测", "独立封存攻击与安全样本 分类别和多语言统计", "独立 QA 签名 原始样本不泄露给研发"],
        ["性能", "1K 8K 32K 128K Token 并发 媒体和故障注入", "冻结拓扑报告含分位数和资源"],
    ], [1.0, 3.65, 2.2])

    add_heading(doc, "11.2 建议能力指标", 2)
    add_body(doc, "下列指标为优化后的建议验收目标，最终数值需由产品、安全、客户环境所有者和独立 QA 冻结。")
    add_table(doc, ["指标", "建议门槛", "说明"], [
        ["总体 Accuracy", "不低于 95%", "沿用测试手册发布门槛"],
        ["总体 Recall", "不低于 95%", "攻击样本总体召回"],
        ["总体 FPR", "不高于 1%", "安全样本误报率"],
        ["总体 FNR", "不高于 5%", "攻击漏报率"],
        ["注入检测 Recall", "不低于 97%", "角色扮演 目标劫持 泄露和间接注入分别统计"],
        ["混淆攻击 Recall", "不低于 95%", "每个编码和字符变体族均达标"],
        ["多轮递进 Recall", "不低于 95%", "至少覆盖 3 至 10 轮攻击链"],
        ["结构化 PII Precision Recall", "均不低于 99%", "对带校验规则的身份证 银行卡等统计"],
        ["输出合规 Recall", "不低于 95%", "政治 色情和保险类别分别出报告"],
        ["变换复检", "100%", "MASK REWRITE SAFE_RESPONSE 后必须复检并通过"],
    ], [2.05, 1.5, 3.3])

    add_heading(doc, "11.3 P0 回归清单", 2)
    add_numbered(doc, [
        "新部署后 policy readiness 返回 ready，并显示可验证 Bundle ID、generation 和公钥摘要。",
        "附件 TC 0028 至 TC 0059 的 32 个维度正反例均返回 HTTP 200，并进入目标维度判定。",
        "正常输入、注入、输出凭证、改写复检、关键检测器超时和流式提交门形成部署 E2E。",
        "签名篡改、公钥不匹配、无绑定和非法状态均拒绝加载；可回滚到最近可用版本。",
        "App 日志不再出现 /app/.next/cache ENOENT；10 个 Worker 无启动错误和异常重启。",
        "重新执行 158 条矩阵，P0 和 P1 失败为 0；外部阻塞项保持真实口径。",
    ])

    add_heading(doc, "12 实施计划与工作包", 1)
    add_table(doc, ["阶段", "周期建议", "工作包", "退出条件"], [
        ["阶段 0", "第 1 周", "密钥分权 策略 bootstrap readiness 缓存整改", "32 条请求不再 POLICY_LOAD_FAILED"],
        ["阶段 1", "第 2 至 3 周", "统一敏感词规则中心 结构化 DLP 输出动作引擎", "词典发布和变换复检 E2E 通过"],
        ["阶段 2", "第 4 至 6 周", "注入语义分类 混淆解码图 多语种与变形测试", "离线分类别指标达到内部门槛"],
        ["阶段 3", "第 7 至 9 周", "图文音视频融合 多轮意图图 算力耗尽防护", "多模态和生命周期预算故障注入通过"],
        ["阶段 4", "第 10 至 11 周", "政治 色情 保险输出策略 模板中心和人工复核", "输出门控和代答全链路通过"],
        ["阶段 5", "第 12 周", "独立盲测 性能 可靠性和交付证据", "具名签收或形成明确外部阻塞清单"],
    ], [0.75, 1.1, 3.45, 1.55])

    add_heading(doc, "12.1 建议职责", 2)
    add_table(doc, ["角色", "职责"], [
        ["技术负责人", "架构边界 策略状态机 API 兼容和发布决策"],
        ["安全算法", "注入 混淆 多语种 多模态和输出合规模型及校准"],
        ["后端平台", "词典 模板 策略包 会话风险 动作执行和审计"],
        ["DevOps", "密钥 注入 readiness 只读容器 资源限额和可观测性"],
        ["行业合规", "政治和保险分类体系 红线词 代答模板和审批"],
        ["QA", "变形测试 盲测 性能 可靠性和 158 条追踪矩阵"],
    ], [1.5, 5.35])

    add_heading(doc, "12.2 代码落点建议", 2)
    add_table(doc, ["代码区域", "建议变更"], [
        ["src/lib/guard-engine-v2/normalization.ts", "拆分解码器注册表 增加解码图预算 切片重组 拼音和扩展混淆骨架"],
        ["src/lib/guard-engine-v2/builtin-detectors.ts", "将静态词表迁移为编译后的签名词典 保留平台 mandatory deny"],
        ["src/lib/guard-engine-v2/semantic-classifier.ts", "增加多模型集成 分类别校准 漂移指标和批次失败隔离"],
        ["src/lib/guard-engine-v2/session-context.ts", "增加会话意图图 风险衰减 状态机和会话级动作"],
        ["src/lib/multimodal 和 src/lib/media", "统一 OCR ASR 视觉 时间轴和跨模态证据格式"],
        ["src/lib/guardrail", "合并 PII 脱敏 改写和安全代答为统一 transformation service"],
        ["src/lib/policy-bundle", "纳入词典 模板 校准和行业策略摘要 增加启动完整性校验"],
        ["docker-compose.yml 和 deploy", "公私钥分权 policy readiness cache tmpfs 和 Secret 门禁"],
    ], [2.65, 4.2])

    add_heading(doc, "13 风险 回滚与运营机制", 1)
    add_table(doc, ["风险", "影响", "控制措施"], [
        ["词典扩张导致误报", "正常业务被阻断", "SHADOW 分类别 FPR 上限 上下文角色和灰度回滚"],
        ["语义模型漂移", "漏报或动作不稳定", "模型摘要固定 周期盲测 校准版本和漂移告警"],
        ["多模态成本过高", "延迟与资源失控", "风险自适应采样 生命周期预算 异步和并发限额"],
        ["代答模板引入新风险", "不合规内容二次输出", "变量白名单 签名审批 渲染后复检和固定兜底"],
        ["签名密钥泄露", "恶意策略可被信任", "KMS HSM 最小权限 轮换 吊销和紧急公钥更新"],
        ["规则与模型结论冲突", "动作不可预测", "明确 mandatory deny 动作优先级和 Judge 触发条件"],
        ["外部联调长期阻塞", "无法完成验收", "提前冻结责任人 环境 数据摘要和时间窗 分项签收"],
    ], [1.65, 2.15, 3.05])

    add_heading(doc, "13.1 回滚策略", 2)
    add_bullets(doc, [
        "每次发布以 Bundle generation 为最小回滚单位，词典、模型、阈值和模板必须同步回滚。",
        "CANARY 指标超过 FPR、错误率或延迟阈值时自动停止放量并恢复 last known good。",
        "数据库迁移采用向后兼容的扩展字段和双读阶段，删除旧字段必须延后到稳定版本。",
        "紧急模式可禁用非关键语义或多模态辅助检测，但平台红线、DLP、输出提交门和策略验签不可关闭。",
    ])

    add_heading(doc, "13.2 运营看板", 2)
    add_bullets(doc, [
        "按方向、风险类别、租户、应用、语言、模态、策略版本和动作统计请求量与命中率。",
        "展示检测器 P50 P95 P99、超时、降级、资源拒绝、Worker backlog 和策略 generation。",
        "跟踪误报申诉、人工复核一致率、badcase 关闭周期、模板兜底率和变换复检失败率。",
        "对策略包缺失、签名失败、公钥不一致、审计链断裂和 mandatory deny 被绕过设置红色告警。",
    ])

    add_heading(doc, "14 结论", 1)
    add_body(doc, "本次优化不建议以继续堆叠静态关键词作为主路线。GuardLLM 已具备规范化、分层检测、语义分类、多模态、会话记忆和资源预算的代码基础，真正的短板是这些能力尚未通过签名策略包形成可部署闭环，以及敏感词、输出动作和效果评测缺少统一治理。")
    add_body(doc, "实施顺序应保持明确：先完成 P0 策略发布与 readiness，恢复核心检测；再统一敏感词和输出变换底座；随后增强注入、混淆、多语种、多模态、多轮和算力攻击；最后在独立盲测和目标环境中完成效果、性能、兼容和可靠性签收。完成这些工作后，平台才能把代码能力转化为可证明、可运营、可回滚的安全能力。")

    add_heading(doc, "附录 A 设计到测试范围映射", 1)
    add_table(doc, ["设计工作包", "重点覆盖用例", "新增证据"], [
        ["策略可用性和签名", "TC 0016 至 TC 0027 TC 0060 至 TC 0074", "policy readiness Bundle 清单 状态流转和回滚 E2E"],
        ["注入和混淆", "TC 0028 TC 0029 及 Guard 绕过相关用例", "直接 间接 多语种 编码 切片和字符变形集"],
        ["结构化 DLP 和输出脱敏", "TC 0030 TC 0048 TC 0050 TC 0058 TC 0059", "字段级 Precision Recall 变换前后摘要和复检"],
        ["多模态", "TC 0097 至 TC 0108", "OCR ASR 视频时间轴 区域证据和协同攻击 E2E"],
        ["多轮和资源", "TC 0019 TC 0024 TC 0141 至 TC 0144", "会话攻击链 长上下文 生命周期预算和租约恢复"],
        ["输出合规和代答", "16 维输出样本及行业扩展用例", "政治 色情 保险分类别指标 模板审批和流式提交门"],
        ["外部验收", "TC 0155 至 TC 0158", "签名性能报告 国产组合 POC 和 Version B 独立结论"],
    ], [2.0, 2.55, 2.3])

    add_heading(doc, "附录 B 参考基线", 1)
    add_bullets(doc, [
        "GuardLLM 附件全量自动化测试报告 2026-09-04。",
        "GuardLLM 158 条用例执行矩阵 2026-09-04。",
        "GuardLLM 全面操作与测试手册 V1.0 2026-09-03。",
        "GuardLLM 配套测试用例 V1.0 2026-09-03。",
        "当前仓库 guard-engine-v2 policy-bundle multimodal media guardrail resource-control 和相关测试代码。",
    ])

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
